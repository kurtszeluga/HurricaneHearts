from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuals"
LOGO = ROOT / "public" / "hurricane-hearts-logo.jpg"

NAVY = "172033"
BLUE = "1F3A5F"
RED = "B42318"
LIGHT_BLUE = "E8EEF5"
LIGHT_RED = "FDECEC"
LIGHT_GRAY = "F1F5F9"
MID_GRAY = "667085"
WHITE = "FFFFFF"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_document(doc, title, compact=False):
    section = doc.sections[0]
    section.top_margin = Inches(0.72 if compact else 0.78)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72 if compact else 0.82)
    section.right_margin = Inches(0.72 if compact else 0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5 if compact else 10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(4 if compact else 6)
    normal.paragraph_format.line_spacing = 1.08 if compact else 1.14

    for name, size, color, before, after in [
        ("Title", 27 if compact else 30, NAVY, 0, 10),
        ("Subtitle", 13, MID_GRAY, 0, 8),
        ("Heading 1", 16, BLUE, 14, 6),
        ("Heading 2", 12.5, RED, 10, 4),
        ("Heading 3", 10.5, BLUE, 7, 3),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = styles["Normal"]
    r = footer.add_run(f"{title}  |  ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_field(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def cover(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(44)
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(2.0))

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(subtitle)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(4.8)
    shade(cell, LIGHT_BLUE)
    set_cell_margins(cell, 180, 220, 180, 220)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Arlington Ridge Community\nVersion 1.0  |  June 2026")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(140)
    r = p.add_run("For community coordination. For emergencies, call 911.")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(RED)
    doc.add_page_break()


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def para(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    list_num_id = None
    if numbered:
        numbering = doc.part.numbering_part.element
        style_num_id = doc.styles[style].element.pPr.numPr.numId.val
        source_num = numbering.xpath(f'./w:num[@w:numId="{style_num_id}"]')[0]
        abstract_num_id = source_num.abstractNumId.val
        existing_ids = [int(num.get(qn("w:numId"))) for num in numbering.xpath("./w:num")]
        list_num_id = max(existing_ids) + 1
        new_num = OxmlElement("w:num")
        new_num.set(qn("w:numId"), str(list_num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(abstract_num_id))
        new_num.append(abstract)
        level_override = OxmlElement("w:lvlOverride")
        level_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        level_override.append(start_override)
        new_num.append(level_override)
        numbering.append(new_num)
    for item in items:
        p = doc.add_paragraph(style=style)
        if numbered:
            p_pr = p._p.get_or_add_pPr()
            num_pr = p_pr.get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = list_num_id
        p.paragraph_format.space_after = Pt(2)
        if isinstance(item, tuple):
            lead, rest = item
            p.add_run(lead).bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def callout(doc, title, text, kind="info"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, LIGHT_RED if kind == "warning" else LIGHT_BLUE)
    set_cell_margins(cell, 120, 150, 120, 150)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(RED if kind == "warning" else BLUE)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def simple_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cell.paragraphs[0].add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            if row_index % 2:
                shade(cells[i], LIGHT_GRAY)
            cells[i].paragraphs[0].add_run(value)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def contents(doc, entries):
    heading(doc, "Contents")
    simple_table(doc, ["Section", "What You Will Find"], entries, [2.2, 4.5])


def add_terms_note(doc):
    callout(
        doc,
        "Terms & Conditions",
        "The current Terms & Conditions are available from the footer at any time. "
        "When a new version is issued, the app may require you to review and accept it at your next login.",
    )


def user_manual():
    title = "Hurricane Hearts User Manual"
    doc = Document()
    configure_document(doc, title)
    cover(doc, title, "A practical guide for residents and Hurricane Hearts volunteers")
    contents(doc, [
        ("Getting Started", "Request access, log in, and complete first-login steps"),
        ("Using the App", "Navigation, profile, dashboard, requests, and directory"),
        ("Privacy & Help", "Contact-information rules, troubleshooting, and safety"),
    ])

    heading(doc, "1. About Hurricane Hearts")
    para(doc, "Hurricane Hearts helps Arlington Ridge residents coordinate neighbor-to-neighbor assistance before, during, and after a storm. It is intended for community support and does not replace emergency services.")
    callout(doc, "Emergency Reminder", "For life-threatening situations, fire, medical emergencies, or immediate danger, call 911.", "warning")
    simple_table(doc, ["Role", "What the Role Can Do"], [
        ("Resident", "Manage a profile, request assistance, volunteer for eligible requests, and use community resources."),
        ("HH Team Member", "Coordinate Hurricane Hearts activities and view contact information needed for that work."),
        ("Admin / Super Admin", "Manage events, accounts, settings, and system records."),
    ], [1.7, 5.0])

    heading(doc, "2. Request Access and Log In")
    heading(doc, "Request Access", 2)
    bullets(doc, [
        "Select Request Access from the home page.",
        "Choose whether you have an email address. If you have email, a separate User ID is optional. If you do not have email, create a User ID.",
        "Enter your name, phone number, house number, street name, and AR lot number.",
        "Open and scroll through the Terms & Conditions, then accept them.",
        "Submit the request and wait for administrator approval.",
    ], numbered=True)
    para(doc, "Address verification compares the house number and AR lot number with the community address list. Street name is still required because it is needed when coordinating assistance.")
    callout(doc, "If the address does not match", "Choose Edit to correct the form, Cancel to stop, or Submit Anyway to send it for administrator review.")
    heading(doc, "Log In", 2)
    bullets(doc, [
        "Enter your email address or User ID.",
        "Enter your password and select Login.",
        "Use Show Password only when no one else can see your screen.",
    ])
    para(doc, "A pending-access message means your request is still awaiting administrator approval. Contact an administrator if you need password or account assistance.")
    heading(doc, "First Login", 2)
    bullets(doc, [
        "Review and accept the current Terms & Conditions when prompted.",
        "Review your profile and correct any contact or address information.",
        "Select the activities you are willing to help with.",
    ])
    add_terms_note(doc)

    heading(doc, "3. Navigation and Dashboard")
    para(doc, "The main navigation provides access to Dashboard, Requests, Directory, and Documents. Additional pages appear only for users with the appropriate role.")
    simple_table(doc, ["Dashboard Area", "Purpose"], [
        ("Summary counts", "Shows Open, Assigned, Partial, Completed, Cancelled, My Requests, and My List."),
        ("New Open Requests", "Shows assistance currently available for volunteers."),
        ("My List", "Shows requests where you volunteered, with Details and Complete actions."),
        ("My Requests", "Shows requests you created and their current status."),
    ], [2.0, 4.7])

    heading(doc, "4. Manage Your Profile")
    bullets(doc, [
        "Open Edit Profile from the navigation/profile area.",
        "Keep your name, phone, email, house number, street name, and AR lot number current.",
        "Select the categories where you are willing to help.",
        "Save your changes.",
    ], numbered=True)
    callout(doc, "Two different category sets", "Your willingness-to-help choices are separate from HH Team coordination assignments. Only authorized administrators manage HH Team assignments.")

    heading(doc, "5. Request Assistance")
    bullets(doc, [
        "Open Requests and select the option to create a new request.",
        "Choose one or more assistance categories.",
        "Select urgency, enter the number of people needed, and describe the need clearly.",
        "Review the information and submit the request.",
    ], numbered=True)
    para(doc, "New requests can be created only while a Hurricane Hearts event is active.")
    simple_table(doc, ["Storm Phase", "Request Categories"], [
        ("Before the Storm", "Storm Preparations"),
        ("During the Storm", "Adopt-A-Buddy; Phone-A-Friend"),
        ("After the Storm", "Cleanup; Borrow or Donate Supplies; Request a Meal; Grocery Store Driver"),
    ], [1.8, 4.9])
    heading(doc, "Request a Meal", 2)
    para(doc, "When selecting Request a Meal, answer Yes or No to the food-allergy question. If Yes, identify all allergens in the comments. Only HH Team members can see and volunteer for Request a Meal items.")

    heading(doc, "6. Volunteer and Complete Assistance")
    bullets(doc, [
        "Open an eligible request and select Volunteer.",
        "Enter how many people you are committing and add a short comment.",
        "For a meal request, select the designated meal preparation location.",
        "Coordinate with the requestor using the contact information shown after volunteering.",
        "When the work is finished, select Complete from My List or the request details.",
    ], numbered=True)
    callout(doc, "Volunteer rules", "You cannot volunteer for your own request. A request may be partially staffed when more help is still needed.")
    heading(doc, "Request Details and History", 2)
    para(doc, "Open Details to see the request description, status, volunteer commitments, completion or cancellation information, and the dated activity history showing who performed each action.")

    heading(doc, "7. Directory, Documents, and Notifications")
    heading(doc, "Resident Directory", 2)
    para(doc, "Use the Directory to search and sort residents and willingness-to-help categories. For privacy, ordinary residents see names and volunteer categories, but not other residents’ phone numbers, email addresses, or addresses.")
    heading(doc, "Documents", 2)
    para(doc, "Use Documents to open community resources such as emergency information, preparation checklists, shelter information, forms, and event-specific materials.")
    heading(doc, "Notifications", 2)
    para(doc, "The app may notify you about request activity. Outgoing emails are informational and include a notice not to reply because the sending account is not monitored.")

    heading(doc, "8. Privacy and Safety")
    bullets(doc, [
        "HH Team members may view resident contact information when needed for coordination.",
        "After someone volunteers for a request, the requestor and volunteer can see each other’s contact information.",
        "Do not share another resident’s contact information outside the purpose of the request.",
        "Never share your password. Log out on shared devices.",
        "Use sound judgment and do not accept work you cannot perform safely.",
    ])

    doc.add_page_break()
    heading(doc, "9. Status Guide and Troubleshooting")
    simple_table(doc, ["Status", "Meaning"], [
        ("Open", "No help has been committed yet."),
        ("Partial", "Some help is committed, but more is still needed."),
        ("Assigned", "The needed help has been committed."),
        ("Completed", "The assistance has been finished."),
        ("Cancelled", "The request is no longer active."),
    ], [1.5, 5.2])
    simple_table(doc, ["Issue", "What to Do"], [
        ("Cannot log in", "Check the email/User ID and password. Contact an administrator for help."),
        ("Waiting for approval", "Your access request has not yet been approved."),
        ("Cannot create a request", "Confirm that an event is active."),
        ("Cannot see a meal request", "Request a Meal items are restricted to HH Team members."),
        ("Contact details are hidden", "This is the normal privacy setting until needed for a volunteer commitment or HH Team work."),
    ], [2.0, 4.7])
    doc.save(OUT / "Hurricane-Hearts-User-Manual.docx")


def admin_manual():
    title = "Hurricane Hearts Admin Manual"
    doc = Document()
    configure_document(doc, title)
    cover(doc, title, "Administration, event coordination, privacy, and go-live procedures")
    contents(doc, [
        ("Administration", "Roles, event control, users, requests, and documents"),
        ("Communication & Records", "Email tools, history, reports, and address verification"),
        ("Operations", "Privacy, troubleshooting, and go-live checklists"),
    ])

    heading(doc, "1. Administration Roles")
    simple_table(doc, ["Role", "Administrative Scope"], [
        ("HH Team Member", "Coordinates assistance and can view resident contact information. This is not an administrator role."),
        ("Admin", "Manages events, residents, requests, documents, email tools, history, and reports."),
        ("Super Admin", "Has the highest privilege and controls sensitive settings such as admin access, User IDs, HH Team assignments, meal preparation locations, and signup verification."),
    ], [1.7, 5.0])
    callout(doc, "Use least privilege", "Assign Admin, Super Admin, and HH Team access only when the person needs it for their responsibilities.", "warning")

    heading(doc, "2. Routine Admin Workflow")
    bullets(doc, [
        "Review pending access requests and address-verification exceptions.",
        "Confirm that user roles, HH Team assignments, and contact information are accurate.",
        "Activate an event only when coordination should begin.",
        "Monitor requests, volunteer commitments, notifications, and email activity during the event.",
        "Deactivate the event after incomplete requests have been reviewed.",
        "Use History and Reports for follow-up and records.",
    ], numbered=True)

    heading(doc, "3. Event Control")
    heading(doc, "Activate an Event", 2)
    bullets(doc, [
        "Open Admin Panel and locate Event Control.",
        "Enter the event name, event date, and comments/information.",
        "Choose whether to send the event email to active users.",
        "Review the confirmation and activate the event.",
    ], numbered=True)
    para(doc, "The event comments/information are included in the new-event email. Event actions are recorded in Event History.")
    heading(doc, "Deactivate an Event", 2)
    bullets(doc, [
        "Review open, partial, and assigned requests.",
        "Resolve or document incomplete work.",
        "Select Deactivate and confirm the action.",
    ], numbered=True)
    callout(doc, "Before deactivation", "The app warns when open or incomplete requests remain. Review them before ending the event.", "warning")

    heading(doc, "4. Admin User Management")
    para(doc, "User Management is the primary account workspace. It supports search, filters, sortable columns, summary counts, and 25-row pagination.")
    heading(doc, "Approve an Access Request", 2)
    bullets(doc, [
        "Find the pending account.",
        "Review the resident’s name, contact information, address, and AR lot number.",
        "If the address failed verification, open Edit and review the highlighted status before approval.",
        "For an override, confirm the approval and enter a short review comment.",
        "Approve the account.",
    ], numbered=True)
    para(doc, "The failed-verification and override information remains in the profile for future reference.")
    heading(doc, "Add a User", 2)
    bullets(doc, [
        "Select Add User and enter the required profile information.",
        "Create a temporary password and use Show Password only when needed.",
        "Save the user and provide the login details securely.",
        "The user must review and accept the Terms & Conditions at first login.",
    ], numbered=True)
    heading(doc, "Edit, Reset, or Delete an Account", 2)
    bullets(doc, [
        ("Edit profile: ", "Update contact/address information and volunteer categories."),
        ("Reset password: ", "A Super Admin can enter a new masked password for the user."),
        ("Activate/deactivate: ", "Control whether the user may access the system."),
        ("Delete: ", "Open Edit, select Delete, confirm, and enter the required reason. Deletion removes both the profile and Firebase Authentication login."),
    ])
    callout(doc, "Sensitive controls", "Only a Super Admin can manage admin access, User IDs, HH Team assignments, managed categories, and meal preparation locations.", "warning")

    heading(doc, "5. HH Team and Meal Coordination")
    para(doc, "HH Team coordination categories are separate from the resident’s willingness-to-help categories.")
    bullets(doc, [
        "A Super Admin designates a profile as an HH Team Member.",
        "A Super Admin selects the categories that the team member coordinates.",
        "A Super Admin may designate a profile address as a Meal Preparation Location.",
        "When an HH Team member volunteers for a Request a Meal item, the volunteer selects a designated preparation location from the available list.",
    ])

    heading(doc, "6. Request Oversight")
    bullets(doc, [
        "Use Dashboard and Requests to monitor open, partial, assigned, completed, and cancelled work.",
        "Open Details to review the full request history, volunteer commitments, comments, and dates.",
        "Admins may edit, cancel, complete, or otherwise manage requests when intervention is needed.",
        "Request a Meal items are restricted to HH Team members.",
        "Contact information is not stored directly in request records; access is controlled through profile permissions and volunteer relationships.",
    ])

    heading(doc, "7. Documents")
    para(doc, "Admins can add, edit, and delete links in the Documents library. Choose an appropriate category and clearly label the resource. Event-specific documents can be associated with an active event.")
    simple_table(doc, ["Document Category Examples", "Typical Content"], [
        ("Emergency Info / Shelter Info", "Official contacts, evacuation, and shelter resources"),
        ("Preparation Checklist / Supplies", "Storm preparation and supply guidance"),
        ("Community Guide / Forms", "Community procedures and forms"),
        ("Event Specific / Other", "Materials relevant to a particular event or special purpose"),
    ], [2.4, 4.3])

    heading(doc, "8. Email and Notifications")
    simple_table(doc, ["Tool", "Use"], [
        ("Email Settings", "Enable or disable automatic delivery and configure recipients for pending-user notices."),
        ("Blast Email", "Send a one-time message to active approved users after confirming the audience and content."),
        ("Recent Email Activity", "Review the delivery audit trail."),
        ("App Notifications", "Review system activity relevant to requests and users."),
    ], [2.0, 4.7])
    callout(doc, "Outgoing messages", "Emails include a notice that recipients should not reply because the sending account is not monitored.")

    heading(doc, "9. History and Reports")
    bullets(doc, [
        "Use History to review event actions and request activity timelines.",
        "Use Reports to export Users, Requests, and Request History as CSV files.",
        "Choose the relevant event or all events, then confirm before downloading.",
        "Store exported files securely because they may contain personal or operational information.",
    ])

    heading(doc, "10. Signup Settings and Address Directory")
    para(doc, "Signup Settings are at the bottom of the Admin Panel and are available only to Super Admins.")
    bullets(doc, [
        "Turn signup address/lot verification on or off.",
        "Upload the approved resident-address CSV.",
        "The CSV must contain: houseNumber, streetName, arLotNumber.",
        "Verification matches house number and AR lot number. Street name remains required and is stored for assistance coordination, but it is not used in the verification match.",
        "A blank row with no house number ends the import; extra unused columns are ignored.",
    ])
    callout(doc, "Verification override", "When a signup does not match, the resident may submit it for admin review. Approval requires an intentional review and comment.", "warning")

    heading(doc, "11. Privacy and Security")
    bullets(doc, [
        "Ordinary residents cannot browse other residents’ email, phone, or address information.",
        "HH Team members can view contact information for coordination responsibilities.",
        "After volunteering, the requestor and volunteer can see each other’s contact information.",
        "Admin Panel access to contact information is for account administration only.",
        "Do not share passwords or export files unnecessarily. Remove access promptly when responsibilities change.",
        "Review roles and permissions periodically.",
    ])

    heading(doc, "12. Before Go-Live Checklist")
    bullets(doc, [
        "Purge test users, requests, history, notifications, and test email activity as appropriate.",
        "Upload the final Arlington Ridge address and lot-number directory.",
        "Enable signup verification when the address list is complete.",
        "Verify Super Admin, Admin, HH Team, managed-category, and meal-preparation-location assignments.",
        "Confirm automatic email settings and pending-user notice recipients.",
        "Test representative Resident, HH Team, Admin, and Super Admin accounts.",
        "Confirm current Terms & Conditions behavior and links.",
        "Confirm production rules and deployment configuration before inviting residents.",
    ])

    heading(doc, "13. Troubleshooting")
    simple_table(doc, ["Issue", "Admin Response"], [
        ("User cannot log in", "Confirm login identifier, active/approved status, Terms acceptance, and reset password if needed."),
        ("User remains pending", "Review and approve the profile; investigate any address mismatch."),
        ("Request cannot be created", "Confirm an event is active and the user is active/approved."),
        ("Meal request is not visible", "Confirm the viewer is designated as an HH Team member."),
        ("Address upload fails", "Confirm headers and required values; remove or leave blank rows after the final address."),
        ("Email not delivered", "Check Email Settings and Recent Email Activity."),
    ], [2.0, 4.7])
    doc.save(OUT / "Hurricane-Hearts-Admin-Manual.docx")


def user_quick():
    title = "Hurricane Hearts User Quick Reference"
    doc = Document()
    configure_document(doc, title, compact=True)
    cover(doc, title, "Common resident tasks at a glance")
    heading(doc, "Login and Access")
    bullets(doc, [
        ("Request access: ", "Use email or create a User ID, enter your profile/address, review the Terms & Conditions, and wait for approval."),
        ("Log in: ", "Use your email or User ID plus password."),
        ("First login: ", "Accept the current Terms & Conditions and review your profile."),
    ])
    heading(doc, "Create a Request")
    bullets(doc, [
        "Open Requests and create a new request while an event is active.",
        "Choose categories, urgency, people needed, and add a clear description.",
        "For Request a Meal, answer the allergy question and list allergens if Yes.",
        "Submit and follow progress in My Requests.",
    ], numbered=True)
    heading(doc, "Volunteer and Complete Help")
    bullets(doc, [
        "Open an eligible request and select Volunteer.",
        "Enter the number of people and a short comment.",
        "Coordinate with the requestor using contact information shown after volunteering.",
        "Select Complete when the assistance is finished.",
    ], numbered=True)
    simple_table(doc, ["Status", "Meaning"], [
        ("Open", "No help committed"),
        ("Partial", "More help needed"),
        ("Assigned", "Needed help committed"),
        ("Completed", "Work finished"),
        ("Cancelled", "No longer active"),
    ], [1.4, 5.3])
    heading(doc, "Privacy and Safety")
    bullets(doc, [
        "Residents see names and volunteer categories, not general contact details.",
        "Requestors and volunteers see each other’s contact information after volunteering.",
        "Do not share passwords or another resident’s contact information.",
        "For emergencies, call 911.",
    ])
    heading(doc, "Where to Find It")
    para(doc, "Dashboard: activity and your work  |  Requests: request, volunteer, or review help  |  "
              "Directory: residents and volunteer categories  |  Documents: community resources  |  "
              "Edit Profile: update your information")
    doc.save(OUT / "Hurricane-Hearts-User-Quick-Reference.docx")


def admin_quick():
    title = "Hurricane Hearts Admin Quick Reference"
    doc = Document()
    configure_document(doc, title, compact=True)
    cover(doc, title, "Common administration tasks at a glance")
    heading(doc, "Start and End an Event")
    simple_table(doc, ["Task", "Steps"], [
        ("Activate", "Admin Panel > Event Control > enter name/date/comments > choose email option > confirm."),
        ("Monitor", "Review Dashboard, Requests, Notifications, and Recent Email Activity."),
        ("Deactivate", "Review incomplete requests > Deactivate > confirm."),
    ], [1.4, 5.3])
    heading(doc, "Review and Manage Users")
    bullets(doc, [
        ("Approve: ", "Review profile and address. For an address mismatch, open Edit, intentionally override, and enter a comment."),
        ("Add: ", "Create the profile and temporary password. The user accepts Terms & Conditions at first login."),
        ("Edit: ", "Update contact/address information and volunteer categories."),
        ("Delete: ", "Open Edit > Delete > confirm > enter reason."),
        ("Super Admin only: ", "Admin access, User IDs, HH Team, managed categories, meal preparation locations, password resets, and Signup Settings."),
    ])
    heading(doc, "Operational Tools")
    simple_table(doc, ["Tool", "Purpose"], [
        ("Requests", "Review details/history and intervene when needed"),
        ("Documents", "Add/edit/delete community resource links"),
        ("Email Settings", "Control automatic email and pending-user recipients"),
        ("Blast Email", "Send confirmed one-time message to active users"),
        ("History", "Review event and request activity"),
        ("Reports", "Export Users, Requests, or Request History CSV"),
        ("Signup Settings", "Toggle verification and upload address CSV"),
    ], [1.8, 4.9])
    doc.add_page_break()
    heading(doc, "Address CSV")
    para(doc, "Required headers: houseNumber, streetName, arLotNumber. Verification matches house number + AR lot number; street is required and stored but not compared.")
    heading(doc, "Privacy Rules")
    bullets(doc, [
        "Residents cannot browse general contact information.",
        "HH Team members can view contact information for coordination.",
        "Requestor and volunteer can see each other after volunteering.",
        "Keep exported files and account access secure.",
    ])
    heading(doc, "Before Go-Live / Start of Event")
    bullets(doc, [
        "Confirm address directory and verification setting.",
        "Verify roles, HH Team categories, and meal preparation locations.",
        "Confirm email settings and recipients.",
        "Test Resident, HH Team, Admin, and Super Admin access.",
        "Review incomplete requests before event deactivation.",
    ])
    callout(doc, "Emergency Reminder", "Hurricane Hearts coordinates community help. For emergencies, call 911.", "warning")
    doc.save(OUT / "Hurricane-Hearts-Admin-Quick-Reference.docx")


if __name__ == "__main__":
    user_manual()
    admin_manual()
    user_quick()
    admin_quick()
    print("Created four Hurricane Hearts Word documents.")

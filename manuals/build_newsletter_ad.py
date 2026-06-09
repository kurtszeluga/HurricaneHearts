from pathlib import Path

import qrcode
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuals"
LOGO = Path("/Users/kurtszeluga/Documents/Hurricane Hearts/Logos/Hurricane Hearts Logo (long).jpg")
QR_PATH = OUT / "hurricane-hearts-registration-qr.png"
DOCX_PATH = OUT / "Hurricane-Hearts-Newsletter-Advertisement.docx"
URL = "https://hurricanehearts.org"

NAVY = "172033"
BLUE = "1F3A5F"
RED = "C51F2B"
LIGHT_BLUE = "EAF0F7"
WHITE = "FFFFFF"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BLUE, size="18"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def make_qr():
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=18,
        border=4,
    )
    qr.add_data(URL)
    qr.make(fit=True)
    image = qr.make_image(fill_color=f"#{NAVY}", back_color=f"#{WHITE}")
    image.save(QR_PATH)


def add_centered_text(cell, text, size, color=NAVY, bold=False, before=0, after=0):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def build():
    make_qr()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(5.5)
    section.top_margin = Inches(0.28)
    section.bottom_margin = Inches(0.28)
    section.left_margin = Inches(0.34)
    section.right_margin = Inches(0.34)
    section.header_distance = Inches(0.1)
    section.footer_distance = Inches(0.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(5.35)
    right.width = Inches(2.35)
    for cell in (left, right):
        set_cell_margins(cell, 130, 180, 130, 180)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    left_p = left.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left_p.paragraph_format.space_after = Pt(2)
    left_p.add_run().add_picture(str(LOGO), width=Inches(3.85))

    add_centered_text(left, "NEIGHBORS HELPING NEIGHBORS", 18, RED, True, after=1)
    add_centered_text(left, "Before, During, and After the Storm", 12.5, BLUE, True, after=7)

    p = left.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        "Hurricane Hearts is Arlington Ridge's resident support network. "
        "Request assistance, volunteer to help, and coordinate community support "
        "for storm preparation, check-ins, cleanup, supplies, meals, and grocery needs."
    )
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    banner = left.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.autofit = False
    banner.cell(0, 0).width = Inches(4.65)
    shade(banner.cell(0, 0), RED)
    set_cell_margins(banner.cell(0, 0), 90, 120, 90, 120)
    bp = banner.cell(0, 0).paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bp.add_run("SIGN UP FOR YOUR RESIDENT ACCOUNT TODAY")
    br.font.name = "Calibri"
    br.font.size = Pt(12)
    br.font.bold = True
    br.font.color.rgb = RGBColor.from_string(WHITE)

    add_centered_text(left, "Access is limited to Arlington Ridge residents and requires approval.", 8.5, BLUE, False, before=6)

    shade(right, LIGHT_BLUE)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp.paragraph_format.space_after = Pt(3)
    rp.add_run().add_picture(str(QR_PATH), width=Inches(1.72))
    add_centered_text(right, "SCAN TO REGISTER", 12, RED, True, after=4)
    add_centered_text(right, "hurricanehearts.org", 9.5, BLUE, True, after=7)
    add_centered_text(right, "Open the website and select\nRequest Access.", 9.5, NAVY, False, after=8)
    add_centered_text(right, "For emergencies, call 911.", 8.5, RED, True)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
